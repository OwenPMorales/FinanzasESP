"""
Genera docs/Entregable_v01_FinanzasESP.docx con el MISMO contenido que el PDF.

Recorre los flowables que arma generar_pdf.build_story() y los traduce a Word:
párrafos con su formato (<b>, <i>, <font face='Consolas'>), títulos con los
estilos Heading de Word (aparecen en el panel de navegación), tablas, bloques de
código y diagramas (rasterizados a PNG).

Uso:  python docs/generar_docx.py
"""
import html
import re
import sys
import tempfile
from pathlib import Path

import pypdfium2
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.graphics import renderPDF
from reportlab.graphics.shapes import Drawing
from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, Preformatted, Spacer, Table

sys.path.insert(0, str(Path(__file__).resolve().parent))
import generar_pdf  # noqa: E402

OUT = generar_pdf.DOCS / "Entregable_v01_FinanzasESP.docx"

TEAL = RGBColor(0x0F, 0x76, 0x6E)
SLATE = RGBColor(0x33, 0x41, 0x55)
GRAY = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Estilo reportlab -> (estilo de Word, tamaño, color, alineación)
PARAGRAPH_STYLES = {
    "t": ("Title", 28, TEAL, WD_ALIGN_PARAGRAPH.CENTER),
    "st": ("Subtitle", 14, SLATE, WD_ALIGN_PARAGRAPH.CENTER),
    "c": ("Normal", 11, GRAY, WD_ALIGN_PARAGRAPH.CENTER),
    "b": ("Normal", None, None, WD_ALIGN_PARAGRAPH.JUSTIFY),
    "cap": ("Caption", 9, GRAY, WD_ALIGN_PARAGRAPH.CENTER),
}
HEADINGS = {"h1": 1, "h2": 2, "h3": 3}


# ----------------------------------------------------------------------------
# Utilidades XML
# ----------------------------------------------------------------------------

def shade(element, hex_color: str) -> None:
    """Fondo de celda o párrafo (w:shd con tipo clear)."""
    props = element.get_or_add_tcPr() if element.tag == qn("w:tc") else element.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    props.append(shd)


def set_font(run, name: str) -> None:
    run.font.name = name
    run.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def force_style_font(style, name: str) -> None:
    """Los estilos del tema usan fuentes del tema (asciiTheme) que ganan sobre font.name."""
    style.font.name = name
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    fonts.set(qn("w:eastAsia"), name)


def add_field(paragraph, instruction: str) -> None:
    """Inserta un campo de Word (ej. PAGE)."""
    for kind, text in (("begin", None), (None, instruction), ("end", None)):
        run = paragraph.add_run()
        if kind:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), kind)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {text} "
            run._r.append(instr)


# ----------------------------------------------------------------------------
# Marcado de reportlab -> runs de Word
# ----------------------------------------------------------------------------

TOKEN = re.compile(r"(<[^>]+>)")


def add_markup(paragraph, markup: str, size=None, color=None, bold=False) -> None:
    """Traduce <b>, <i>, <font face=...> y <br/> a runs con formato."""
    state = {"b": bold, "i": False, "mono": False}
    for token in TOKEN.split(markup):
        if not token:
            continue
        tag = token.lower()
        if tag.startswith("<"):
            if tag in ("<b>", "<strong>"):
                state["b"] = True
            elif tag in ("</b>", "</strong>"):
                state["b"] = bold
            elif tag == "<i>":
                state["i"] = True
            elif tag == "</i>":
                state["i"] = False
            elif tag.startswith("<font") and "consolas" in tag:
                state["mono"] = True
            elif tag == "</font>":
                state["mono"] = False
            elif tag.startswith("<br"):
                paragraph.add_run().add_break()
            continue
        run = paragraph.add_run(html.unescape(token).replace("\xa0", " "))
        run.bold = state["b"] or None
        run.italic = state["i"] or None
        if state["mono"]:
            set_font(run, "Consolas")
            run.font.size = Pt((size or 10.5) - 1)
        elif size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color


# ----------------------------------------------------------------------------
# Flowables -> Word
# ----------------------------------------------------------------------------

class Converter:
    def __init__(self) -> None:
        self.doc = Document()
        self.tmp = Path(tempfile.mkdtemp(prefix="finanzas_docx_"))
        self.figures = 0
        self.pending_space = 0.0
        self._setup_document()

    def _setup_document(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.59), Cm(27.94)  # carta
        for side in ("left_margin", "right_margin", "top_margin"):
            setattr(section, side, Cm(2))
        section.bottom_margin = Cm(2.2)

        styles = self.doc.styles
        normal = styles["Normal"]
        force_style_font(normal, "Arial")
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.2

        for level, (size, color) in {1: (18, TEAL), 2: (13.5, SLATE), 3: (11, SLATE)}.items():
            style = styles[f"Heading {level}"]
            force_style_font(style, "Arial")
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.italic = False
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(14 if level < 3 else 8)
            style.paragraph_format.space_after = Pt(6)

        for name in ("Title", "Subtitle", "Caption"):
            style = styles[name]
            force_style_font(style, "Arial")
            style.font.italic = name == "Caption"
            # El tema de Word pone borde inferior y espaciado de letras al Title
            pPr = style.element.get_or_add_pPr()
            for border in pPr.findall(qn("w:pBdr")):
                pPr.remove(border)

        core = self.doc.core_properties
        core.title = "FinanzasESP — Entregable v01"
        core.author = "Owen Morales"
        core.subject = "Modelo User, Hasheo, Token y MinIO"

    def _footer(self) -> None:
        """Pie de página desde la página 2 (la portada va sin pie, como en el PDF)."""
        first = self.doc.sections[0]
        first.different_first_page_header_footer = True
        footer = first.footer.paragraphs[0]
        footer.text = ""
        add_markup(footer, "FinanzasESP — Entregable v01: User, Hasheo, Token y MinIO", size=8.5, color=GRAY)
        footer.add_run("\t")
        page = footer.add_run("Página ")
        page.font.size, page.font.color.rgb = Pt(8.5), GRAY
        add_field(footer, "PAGE")
        for run in footer.runs[-3:]:
            run.font.size, run.font.color.rgb = Pt(8.5), GRAY
        # tabulador derecho al margen
        tabs = footer.paragraph_format.tab_stops
        for stop in list(tabs):
            tabs._pPr.remove(stop._tab)
        tabs.add_tab_stop(Cm(17.59), alignment=2)  # 2 = WD_TAB_ALIGNMENT.RIGHT

    def _apply_space(self, paragraph) -> None:
        if self.pending_space:
            paragraph.paragraph_format.space_before = Pt(min(self.pending_space, 140))
            self.pending_space = 0.0

    # --- flowables ---

    def paragraph(self, flowable: Paragraph) -> None:
        name = flowable.style.name
        text = flowable.text
        if name in HEADINGS:
            p = self.doc.add_heading(level=HEADINGS[name])
            add_markup(p, text)
        elif name == "bl":
            numbered = str(flowable.bulletText).rstrip(".").isdigit()
            p = self.doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            indent = text.count("&nbsp;") // 3
            if indent:
                p.paragraph_format.left_indent = Cm(0.63 + 0.75 * indent)
            add_markup(p, text)
        else:
            style, size, color, align = PARAGRAPH_STYLES.get(name, PARAGRAPH_STYLES["b"])
            p = self.doc.add_paragraph(style=style)
            p.alignment = align
            add_markup(p, text, size=size, color=color, bold=name == "t")
        self._apply_space(p)

    def code(self, flowable: Preformatted) -> None:
        p = self.doc.add_paragraph()
        shade(p._p, "F1F5F9")
        fmt = p.paragraph_format
        fmt.left_indent = fmt.right_indent = Cm(0.2)
        fmt.space_before, fmt.space_after = Pt(4), Pt(10)
        fmt.line_spacing = 1.15  # con 1.0 Word recorta los guiones bajos
        for i, line in enumerate(flowable.lines):
            if i:
                p.runs[-1].add_break()
            run = p.add_run(line or " ")
            set_font(run, "Consolas")
            run.font.size = Pt(8)
        self._apply_space(p)

    def table(self, flowable: Table) -> None:
        rows = flowable._cellvalues
        widths = flowable._colWidths
        table = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                cell.width = Pt(widths[c])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                header = r == 0
                text = value.text if isinstance(value, Paragraph) else str(value)
                add_markup(p, text, size=9, color=WHITE if header else None, bold=header)
                if header:
                    shade(cell._tc, "0F766E")
                elif r % 2 == 0:
                    shade(cell._tc, "F8FAFC")
        # Márgenes internos de las celdas (en twips: 1 pt = 20)
        margins = OxmlElement("w:tblCellMar")
        for side, value in (("top", 60), ("bottom", 60), ("left", 100), ("right", 100)):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), str(value))
            el.set(qn("w:type"), "dxa")
            margins.append(el)
        table._tbl.tblPr.append(margins)
        # La primera fila se repite si la tabla cruza de página
        trPr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        trPr.append(repeat)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def drawing(self, flowable: Drawing) -> None:
        self.figures += 1
        pdf_path = self.tmp / f"figura{self.figures}.pdf"
        png_path = self.tmp / f"figura{self.figures}.png"
        renderPDF.drawToFile(flowable, str(pdf_path))
        pdf = pypdfium2.PdfDocument(str(pdf_path))
        pdf[0].render(scale=4).to_pil().save(png_path)
        pdf.close()
        self._picture(png_path, flowable.width * 0.9)  # Word ocupa algo más de alto que reportlab

    def image(self, flowable: Image) -> None:
        self._picture(Path(flowable.filename), flowable.drawWidth)

    def _picture(self, path: Path, width_pt: float) -> None:
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True  # la figura no se separa de su pie
        p.add_run().add_picture(str(path), width=Pt(min(width_pt, 17.59 / 2.54 * 72)))
        self._apply_space(p)

    def page_break(self) -> None:
        self.pending_space = 0.0
        last = self.doc.paragraphs[-1]
        last.add_run().add_break(WD_BREAK.PAGE)

    def convert(self, story: list) -> None:
        # La plantilla de python-docx puede traer un párrafo vacío al inicio.
        if self.doc.paragraphs and not self.doc.paragraphs[0].text:
            self.doc._body._body.remove(self.doc.paragraphs[0]._p)
        for flowable in story:
            self.dispatch(flowable)
        self._footer()

    def dispatch(self, flowable) -> None:
        if isinstance(flowable, KeepTogether):
            for inner in flowable._content:
                self.dispatch(inner)
        elif isinstance(flowable, Preformatted):
            self.code(flowable)
        elif isinstance(flowable, Paragraph):
            self.paragraph(flowable)
        elif isinstance(flowable, Table):
            self.table(flowable)
        elif isinstance(flowable, Drawing):
            self.drawing(flowable)
        elif isinstance(flowable, Image):
            self.image(flowable)
        elif isinstance(flowable, Spacer):
            self.pending_space += flowable.height
        elif isinstance(flowable, PageBreak):
            self.page_break()
        else:
            raise TypeError(f"Flowable no soportado: {type(flowable).__name__}")

    def save(self) -> None:
        self.doc.save(OUT)
        print(f"DOCX generado: {OUT}")


def build() -> None:
    converter = Converter()
    converter.convert(generar_pdf.build_story())
    converter.save()


if __name__ == "__main__":
    build()
